from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from app.auth import get_current_user
from app.database import get_db
from app.models import (
    Device,
    EmployeeWorkHour,
    Equipment,
    EquipmentMaintenanceOrder,
    EquipmentOeeSnapshot,
    EquipmentOutputRecord,
    EquipmentRepair,
    EquipmentRuntimeLog,
    InspectionRecord,
    InspectionRecordItem,
    ProductionLine,
    ProductionOutputRecord,
    ProductionPlan,
    User,
    WorkOrder,
)
from app.schemas import (
    DailyOutputLinesResponse,
    DailyOutputReportItem,
    EmployeeWorkHourFilterEmployee,
    EmployeeWorkHourFiltersResponse,
    EmployeeWorkHourReportItem,
    EmployeeWorkHourReportListResponse,
    EquipmentDowntimeFilterEquipment,
    EquipmentDowntimeFiltersResponse,
    EquipmentDowntimeDimensionStat,
    EquipmentDowntimeReasonParetoItem,
    EquipmentDowntimeReliabilityMetric,
    EquipmentDowntimeReportItem,
    EquipmentDowntimeReportListResponse,
    EquipmentDowntimeSummary,
    EquipmentDowntimeTrendPoint,
    EquipmentOeeFilterEquipment,
    EquipmentOeeFiltersResponse,
    EquipmentOeeReportItem,
    EquipmentOeeReportListResponse,
    EquipmentOeeSummary,
    EquipmentOeeTrendPoint,
    EquipmentInspectionFilterDevice,
    EquipmentInspectionFiltersResponse,
    EquipmentInspectionReportItem,
    EquipmentInspectionReportListResponse,
    EquipmentMaintenanceReportItem,
    EquipmentMaintenanceReportListResponse,
    EquipmentMaintenanceReportSummary,
    EquipmentRepairDetail,
    EquipmentRepairListItem,
    EquipmentRepairListResponse,
    EquipmentRepairPartResponse,
    WipReportItem,
    WipReportListResponse,
    WipReportProcessesResponse,
)
from app.work_order_utils import (
    STANDARD_PROCESSES,
    WIP_EXCLUDED_STATUSES,
    calc_wip_quantity,
    derive_current_process,
)

router = APIRouter(prefix="/api/reports", tags=["报表中心"])

WIP_METRIC = "wip"


class DailyOutputReportItemOut(DailyOutputReportItem):
    """日产报表响应行（扩展车间、工时、生产人员）。"""

    workshop: str | None = Field(default=None, description="车间")
    work_hours: float = Field(default=0, description="工时（小时）")
    production_staff: str | None = Field(default=None, description="生产人员（逗号分隔）")


class DailyOutputReportListResponseOut(BaseModel):
    items: list[DailyOutputReportItemOut]
    total: int
    page: int
    page_size: int
    plan_qty_sum: int = Field(description="当前筛选条件下计划产量合计")
    actual_qty_sum: int = Field(description="当前筛选条件下实际产量合计")
    defect_qty_sum: int = Field(description="当前筛选条件下不良数量合计")
    work_hours_sum: float = Field(default=0, description="当前筛选条件下工时合计")


class DailyOutputFiltersResponse(BaseModel):
    lines: list[str] = Field(description="可选产线名称列表")
    workshops: list[str] = Field(description="可选车间名称列表")


def _to_wip_item(wo: WorkOrder) -> WipReportItem:
    process = wo.current_process or derive_current_process(
        wo.status, wo.plan_quantity, wo.actual_quantity
    )
    return WipReportItem(
        id=wo.id,
        order_no=wo.order_no,
        product_name=wo.product_name,
        current_process=process,
        wip_quantity=calc_wip_quantity(wo.plan_quantity, wo.actual_quantity),
        status=wo.status,
        start_date=wo.start_date,
        end_date=wo.end_date,
        plan_quantity=wo.plan_quantity,
        actual_quantity=wo.actual_quantity,
    )


@router.get(
    "/wip",
    response_model=WipReportListResponse,
    summary="在制品报表",
    description=(
        "按工单列出在制品数据（wip 口径：未完工且非取消的工单）。"
        "在制数量 = 计划数量 - 实际数量；工序取自 work_orders.current_process，"
        "为空时按完成进度推导（待开工为贴片）。支持按状态、工序、计划日期范围筛选。"
    ),
)
def list_wip_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    status: str | None = Query(None, description="工单状态筛选"),
    process: str | None = Query(None, description="工序筛选"),
    start_date: date | None = Query(None, description="计划开始日期（起）"),
    end_date: date | None = Query(None, description="计划结束日期（止）"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(WorkOrder).filter(~WorkOrder.status.in_(WIP_EXCLUDED_STATUSES))

    if status:
        query = query.filter(WorkOrder.status == status)
    if process:
        query = query.filter(WorkOrder.current_process == process)
    if start_date:
        query = query.filter(WorkOrder.start_date >= start_date)
    if end_date:
        query = query.filter(WorkOrder.end_date <= end_date)

    query = query.order_by(WorkOrder.start_date.desc(), WorkOrder.id.desc())
    total = query.count()
    rows = query.offset((page - 1) * page_size).limit(page_size).all()

    return WipReportListResponse(
        items=[_to_wip_item(wo) for wo in rows],
        total=total,
        page=page,
        page_size=page_size,
        metric=WIP_METRIC,
    )


@router.get(
    "/wip/processes",
    response_model=WipReportProcessesResponse,
    summary="在制品报表工序选项",
    description="返回标准工序列表，供报表筛选下拉使用。",
)
def list_wip_processes(
    _current_user: User = Depends(get_current_user),
):
    return WipReportProcessesResponse(processes=list(STANDARD_PROCESSES))


def _as_date(value) -> date:
    if isinstance(value, date) and not isinstance(value, datetime):
        return value
    if isinstance(value, datetime):
        return value.date()
    return date.fromisoformat(str(value)[:10])


def _resolve_daily_output_line_ids(
    db: Session,
    *,
    production_line: str | None,
    workshop: str | None,
) -> list[int] | None:
    """按产线/车间解析产线 ID；无匹配时返回空列表，未指定筛选时返回 None。"""
    if not production_line and not workshop:
        return None
    query = db.query(ProductionLine)
    if production_line:
        query = query.filter(ProductionLine.name == production_line)
    if workshop:
        query = query.filter(ProductionLine.workshop == workshop)
    line_ids = [ln.id for ln in query.all()]
    return line_ids


def _build_work_hour_staff_maps(
    db: Session,
    *,
    date_from: date,
    date_to: date,
) -> tuple[dict[tuple[date, str], float], dict[tuple[date, str], set[str]], dict[date, float], dict[date, set[str]]]:
    """构建 (日期, 部门/车间) 与日期维度的工时、人员映射。"""
    workshop_hours: dict[tuple[date, str], float] = defaultdict(float)
    workshop_staff: dict[tuple[date, str], set[str]] = defaultdict(set)
    date_hours: dict[date, float] = defaultdict(float)
    date_staff: dict[date, set[str]] = defaultdict(set)

    rows = (
        db.query(EmployeeWorkHour)
        .filter(
            EmployeeWorkHour.work_date >= date_from,
            EmployeeWorkHour.work_date <= date_to,
        )
        .all()
    )
    for row in rows:
        work_date = _as_date(row.work_date)
        hours = float(row.work_hours or 0)
        dept = (row.department or "").strip()
        date_hours[work_date] += hours
        date_staff[work_date].add(row.employee_name)
        if dept:
            key = (work_date, dept)
            workshop_hours[key] += hours
            workshop_staff[key].add(row.employee_name)

    return workshop_hours, workshop_staff, date_hours, date_staff


def _lookup_work_hour_staff(
    report_date: date,
    workshop: str | None,
    *,
    workshop_hours: dict[tuple[date, str], float],
    workshop_staff: dict[tuple[date, str], set[str]],
    date_hours: dict[date, float],
    date_staff: dict[date, set[str]],
) -> tuple[float, str | None]:
    """优先按车间匹配工时记录，否则回退到当日汇总。"""
    hours = 0.0
    names: set[str] = set()
    if workshop:
        hours = workshop_hours.get((report_date, workshop), 0.0)
        names = workshop_staff.get((report_date, workshop), set())
    if not hours and not names:
        hours = date_hours.get(report_date, 0.0)
        names = date_staff.get(report_date, set())
    staff = "、".join(sorted(names)) if names else None
    return round(hours, 2), staff


def _build_daily_output_rows(
    db: Session,
    *,
    date_from: date,
    date_to: date,
    production_line: str | None,
    workshop: str | None,
) -> list[DailyOutputReportItemOut]:
    """按日 / 车间 / 产线聚合产量事实与计划，并关联工时与生产人员。"""
    line_ids = _resolve_daily_output_line_ids(
        db, production_line=production_line, workshop=workshop
    )
    if line_ids is not None and not line_ids:
        return []

    day_start = datetime.combine(date_from, datetime.min.time())
    day_end = datetime.combine(date_to + timedelta(days=1), datetime.min.time())

    out_q = (
        db.query(
            func.date(ProductionOutputRecord.record_at).label("report_date"),
            ProductionOutputRecord.production_line_id,
            func.coalesce(func.sum(ProductionOutputRecord.actual_qty), 0).label("actual_qty"),
            func.coalesce(func.sum(ProductionOutputRecord.defect_qty), 0).label("defect_qty"),
            func.coalesce(func.sum(ProductionOutputRecord.area_output), 0).label("area_output"),
        )
        .filter(
            ProductionOutputRecord.record_at >= day_start,
            ProductionOutputRecord.record_at < day_end,
        )
        .group_by(
            func.date(ProductionOutputRecord.record_at),
            ProductionOutputRecord.production_line_id,
        )
    )
    if line_ids is not None:
        out_q = out_q.filter(ProductionOutputRecord.production_line_id.in_(line_ids))

    output_rows = out_q.all()

    plan_q = (
        db.query(
            ProductionPlan.plan_date,
            ProductionPlan.production_line_id,
            func.coalesce(func.sum(ProductionPlan.plan_qty), 0).label("plan_qty"),
        )
        .filter(
            ProductionPlan.plan_date >= date_from,
            ProductionPlan.plan_date <= date_to,
        )
        .group_by(
            ProductionPlan.plan_date,
            ProductionPlan.production_line_id,
        )
    )
    if line_ids is not None:
        plan_q = plan_q.filter(ProductionPlan.production_line_id.in_(line_ids))

    plan_map: dict[tuple, int] = {}
    for row in plan_q.all():
        plan_map[(_as_date(row.plan_date), row.production_line_id)] = int(row.plan_qty or 0)

    line_rows = db.query(ProductionLine).all()
    if line_ids is not None:
        line_rows = [ln for ln in line_rows if ln.id in line_ids]
    line_map = {ln.id: ln.name for ln in line_rows}
    workshop_map = {ln.id: ln.workshop for ln in line_rows}

    wh_maps = _build_work_hour_staff_maps(db, date_from=date_from, date_to=date_to)
    workshop_hours, workshop_staff, date_hours, date_staff = wh_maps

    keys: set[tuple] = set()
    out_map: dict[tuple, tuple[int, int, float]] = {}
    for row in output_rows:
        report_date = _as_date(row.report_date)
        key = (report_date, row.production_line_id)
        keys.add(key)
        out_map[key] = (int(row.actual_qty or 0), int(row.defect_qty or 0), float(row.area_output or 0))

    for key in plan_map:
        keys.add(key)

    items: list[DailyOutputReportItemOut] = []
    for report_date, line_id in sorted(
        keys, key=lambda k: (k[0], line_map.get(k[1], "")), reverse=True
    ):
        actual_qty, defect_qty, area_output = out_map.get((report_date, line_id), (0, 0, 0.0))
        plan_qty = plan_map.get((report_date, line_id), 0)
        line_workshop = workshop_map.get(line_id)
        work_hours, production_staff = _lookup_work_hour_staff(
            report_date,
            line_workshop,
            workshop_hours=workshop_hours,
            workshop_staff=workshop_staff,
            date_hours=date_hours,
            date_staff=date_staff,
        )
        achievement = round(actual_qty / plan_qty * 100, 2) if plan_qty else 0.0
        defect_rate = (
            round(defect_qty / (actual_qty + defect_qty) * 100, 2)
            if (actual_qty + defect_qty) > 0
            else 0.0
        )
        items.append(
            DailyOutputReportItemOut(
                report_date=report_date,
                production_line=line_map.get(line_id, f"产线#{line_id}"),
                product_code=None,
                product_name=None,
                plan_qty=plan_qty,
                actual_qty=actual_qty,
                defect_qty=defect_qty,
                area_output=round(area_output, 2),
                achievement_rate=achievement,
                defect_rate=defect_rate,
                workshop=line_workshop,
                work_hours=work_hours,
                production_staff=production_staff,
            )
        )
    return items


@router.get(
    "/daily-output",
    response_model=DailyOutputReportListResponseOut,
    summary="日产报表",
    description=(
        "按生产日期、车间、产线聚合日产量："
        "实际/不良来自 production_output_records，计划来自 production_plans，"
        "工时与生产人员来自 employee_work_hours。"
        "默认查询近 7 日（含今天）；支持单日/区间、车间、产线筛选与分页。"
        "报表中心「日产报表」页数据来自本接口。"
    ),
)
def list_daily_output_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    date_from: date | None = Query(None, description="生产日期起（含）"),
    date_to: date | None = Query(None, description="生产日期止（含）"),
    production_line: str | None = Query(None, description="产线名称（精确匹配）"),
    workshop: str | None = Query(None, description="车间名称（精确匹配）"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    today = date.today()
    if date_to is None:
        date_to = today
    if date_from is None:
        date_from = date_to - timedelta(days=6)
    if date_from > date_to:
        date_from, date_to = date_to, date_from

    all_items = _build_daily_output_rows(
        db,
        date_from=date_from,
        date_to=date_to,
        production_line=production_line,
        workshop=workshop,
    )
    total = len(all_items)
    plan_sum = sum(i.plan_qty for i in all_items)
    actual_sum = sum(i.actual_qty for i in all_items)
    defect_sum = sum(i.defect_qty for i in all_items)
    work_hours_sum = round(sum(i.work_hours for i in all_items), 2)
    start = (page - 1) * page_size
    page_items = all_items[start : start + page_size]

    return DailyOutputReportListResponseOut(
        items=page_items,
        total=total,
        page=page,
        page_size=page_size,
        plan_qty_sum=plan_sum,
        actual_qty_sum=actual_sum,
        defect_qty_sum=defect_sum,
        work_hours_sum=work_hours_sum,
    )


@router.get(
    "/daily-output/filters",
    response_model=DailyOutputFiltersResponse,
    summary="日产报表筛选选项",
    description="返回产线、车间名称列表，供日产报表筛选下拉使用。",
)
def list_daily_output_filters(
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    lines = [
        name
        for (name,) in db.query(ProductionLine.name).order_by(ProductionLine.id).all()
    ]
    workshops = sorted(
        {
            ws
            for (ws,) in db.query(ProductionLine.workshop)
            .filter(ProductionLine.workshop.isnot(None), ProductionLine.workshop != "")
            .distinct()
            .all()
            if ws
        }
    )
    return DailyOutputFiltersResponse(lines=lines, workshops=workshops)


@router.get(
    "/daily-output/lines",
    response_model=DailyOutputLinesResponse,
    summary="日产报表产线选项",
    description="返回产线名称列表，供日产报表筛选下拉使用（兼容旧客户端）。",
)
def list_daily_output_lines(
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    lines = [
        name
        for (name,) in db.query(ProductionLine.name).order_by(ProductionLine.id).all()
    ]
    return DailyOutputLinesResponse(lines=lines)


WORK_HOUR_DIMENSIONS = ("detail", "employee", "employee_date", "employee_month", "project", "department")

APPROVAL_STATUS_LABELS = {
    "pending": "待审批",
    "approved": "已通过",
    "rejected": "已驳回",
}


def _approval_label(status: str | None) -> str:
    if not status:
        return "—"
    return APPROVAL_STATUS_LABELS.get(status, status)


def _apply_work_hour_filters(
    query,
    *,
    date_from: date,
    date_to: date,
    department: str | None,
    employee_no: str | None,
    project_name: str | None,
):
    query = query.filter(
        EmployeeWorkHour.work_date >= date_from,
        EmployeeWorkHour.work_date <= date_to,
    )
    if department:
        query = query.filter(EmployeeWorkHour.department == department)
    if employee_no:
        query = query.filter(EmployeeWorkHour.employee_no == employee_no)
    if project_name:
        query = query.filter(EmployeeWorkHour.project_name == project_name)
    return query


def _build_work_hour_report_items(
    db: Session,
    *,
    date_from: date,
    date_to: date,
    department: str | None,
    employee_no: str | None,
    project_name: str | None,
    dimension: str,
) -> list[EmployeeWorkHourReportItem]:
    query = db.query(EmployeeWorkHour)
    query = _apply_work_hour_filters(
        query,
        date_from=date_from,
        date_to=date_to,
        department=department,
        employee_no=employee_no,
        project_name=project_name,
    )

    if dimension == "detail":
        rows = query.order_by(
            EmployeeWorkHour.work_date.desc(),
            EmployeeWorkHour.employee_no,
            EmployeeWorkHour.id.desc(),
        ).all()
        return [
            EmployeeWorkHourReportItem(
                employee_name=row.employee_name,
                employee_no=row.employee_no,
                department=row.department,
                project_name=row.project_name,
                task_name=row.task_name,
                work_date=row.work_date,
                work_hours=float(row.work_hours or 0),
                overtime_hours=float(row.overtime_hours or 0),
                approval_status=_approval_label(row.approval_status),
            )
            for row in rows
        ]


    if dimension == "employee":
        rows = (
            query.with_entities(
                EmployeeWorkHour.employee_no,
                EmployeeWorkHour.employee_name,
                EmployeeWorkHour.department,
                func.sum(EmployeeWorkHour.work_hours).label("work_hours"),
                func.sum(EmployeeWorkHour.overtime_hours).label("overtime_hours"),
                func.count(EmployeeWorkHour.id).label("record_count"),
            )
            .group_by(
                EmployeeWorkHour.employee_no,
                EmployeeWorkHour.employee_name,
                EmployeeWorkHour.department,
            )
            .order_by(EmployeeWorkHour.employee_no)
            .all()
        )
        return [
            EmployeeWorkHourReportItem(
                employee_name=row.employee_name,
                employee_no=row.employee_no,
                department=row.department,
                work_hours=round(float(row.work_hours or 0), 2),
                overtime_hours=round(float(row.overtime_hours or 0), 2),
                approval_status="汇总",
                record_count=int(row.record_count or 0),
            )
            for row in rows
        ]

    if dimension == "employee_date":
        rows = (
            query.with_entities(
                EmployeeWorkHour.employee_no,
                EmployeeWorkHour.employee_name,
                EmployeeWorkHour.department,
                EmployeeWorkHour.work_date,
                func.sum(EmployeeWorkHour.work_hours).label("work_hours"),
                func.sum(EmployeeWorkHour.overtime_hours).label("overtime_hours"),
                func.count(EmployeeWorkHour.id).label("record_count"),
            )
            .group_by(
                EmployeeWorkHour.employee_no,
                EmployeeWorkHour.employee_name,
                EmployeeWorkHour.department,
                EmployeeWorkHour.work_date,
            )
            .order_by(EmployeeWorkHour.work_date.desc(), EmployeeWorkHour.employee_no)
            .all()
        )
        return [
            EmployeeWorkHourReportItem(
                employee_name=row.employee_name,
                employee_no=row.employee_no,
                department=row.department,
                work_date=row.work_date,
                work_hours=round(float(row.work_hours or 0), 2),
                overtime_hours=round(float(row.overtime_hours or 0), 2),
                approval_status="汇总",
                record_count=int(row.record_count or 0),
            )
            for row in rows
        ]

    if dimension == "employee_month":
        month_expr = func.strftime("%Y-%m", EmployeeWorkHour.work_date)
        rows = (
            query.with_entities(
                EmployeeWorkHour.employee_no,
                EmployeeWorkHour.employee_name,
                EmployeeWorkHour.department,
                month_expr.label("work_month"),
                func.sum(EmployeeWorkHour.work_hours).label("work_hours"),
                func.sum(EmployeeWorkHour.overtime_hours).label("overtime_hours"),
                func.count(EmployeeWorkHour.id).label("record_count"),
            )
            .group_by(
                EmployeeWorkHour.employee_no,
                EmployeeWorkHour.employee_name,
                EmployeeWorkHour.department,
                month_expr,
            )
            .order_by(month_expr.desc(), EmployeeWorkHour.employee_no)
            .all()
        )
        return [
            EmployeeWorkHourReportItem(
                employee_name=row.employee_name,
                employee_no=row.employee_no,
                department=row.department,
                work_month=row.work_month,
                work_hours=round(float(row.work_hours or 0), 2),
                overtime_hours=round(float(row.overtime_hours or 0), 2),
                approval_status="汇总",
                record_count=int(row.record_count or 0),
            )
            for row in rows
        ]

    if dimension == "project":
        rows = (
            query.with_entities(
                EmployeeWorkHour.project_name,
                func.sum(EmployeeWorkHour.work_hours).label("work_hours"),
                func.sum(EmployeeWorkHour.overtime_hours).label("overtime_hours"),
                func.count(EmployeeWorkHour.id).label("record_count"),
            )
            .group_by(EmployeeWorkHour.project_name)
            .order_by(EmployeeWorkHour.project_name)
            .all()
        )
        return [
            EmployeeWorkHourReportItem(
                employee_name="—",
                employee_no="—",
                department="—",
                project_name=row.project_name,
                work_hours=round(float(row.work_hours or 0), 2),
                overtime_hours=round(float(row.overtime_hours or 0), 2),
                approval_status="汇总",
                record_count=int(row.record_count or 0),
            )
            for row in rows
        ]

    # department
    rows = (
        query.with_entities(
            EmployeeWorkHour.department,
            func.sum(EmployeeWorkHour.work_hours).label("work_hours"),
            func.sum(EmployeeWorkHour.overtime_hours).label("overtime_hours"),
            func.count(EmployeeWorkHour.id).label("record_count"),
        )
        .group_by(EmployeeWorkHour.department)
        .order_by(EmployeeWorkHour.department)
        .all()
    )
    return [
        EmployeeWorkHourReportItem(
            employee_name="—",
            employee_no="—",
            department=row.department,
            work_hours=round(float(row.work_hours or 0), 2),
            overtime_hours=round(float(row.overtime_hours or 0), 2),
            approval_status="汇总",
            record_count=int(row.record_count or 0),
        )
        for row in rows
    ]


def _normalize_work_hour_date_range(
    date_from: date | None, date_to: date | None
) -> tuple[date, date]:
    today = date.today()
    if date_to is None:
        date_to = today
    if date_from is None:
        date_from = date_to - timedelta(days=29)
    if date_from > date_to:
        date_from, date_to = date_to, date_from
    return date_from, date_to


def _work_hour_export_headers(dimension: str) -> list[str]:
    if dimension == "employee":
        return [
            "员工姓名",
            "工号",
            "所属部门",
            "工时数",
            "加班工时",
            "明细条数",
            "审批/状态",
        ]
    if dimension == "detail":
        return [
            "员工姓名",
            "工号",
            "所属部门",
            "项目名称",
            "任务名称",
            "日期",
            "工时数",
            "加班工时",
            "审批/状态",
        ]
    if dimension == "employee_date":
        return [
            "员工姓名",
            "工号",
            "所属部门",
            "日期",
            "工时数",
            "加班工时",
            "明细条数",
            "审批/状态",
        ]
    if dimension == "employee_month":
        return [
            "员工姓名",
            "工号",
            "所属部门",
            "月份",
            "工时数",
            "加班工时",
            "明细条数",
            "审批/状态",
        ]
    if dimension == "project":
        return ["项目名称", "工时数", "加班工时", "明细条数", "审批/状态"]
    return ["所属部门", "工时数", "加班工时", "明细条数", "审批/状态"]


def _work_hour_export_row(item: EmployeeWorkHourReportItem, dimension: str) -> list:
    if dimension == "employee":
        return [
            item.employee_name,
            item.employee_no,
            item.department,
            item.work_hours,
            item.overtime_hours,
            item.record_count or 0,
            item.approval_status or "",
        ]
    if dimension == "detail":
        return [
            item.employee_name,
            item.employee_no,
            item.department,
            item.project_name or "",
            item.task_name or "",
            item.work_date.isoformat() if item.work_date else "",
            item.work_hours,
            item.overtime_hours,
            item.approval_status or "",
        ]
    if dimension == "employee_date":
        return [
            item.employee_name,
            item.employee_no,
            item.department,
            item.work_date.isoformat() if item.work_date else "",
            item.work_hours,
            item.overtime_hours,
            item.record_count or 0,
            item.approval_status or "",
        ]
    if dimension == "employee_month":
        return [
            item.employee_name,
            item.employee_no,
            item.department,
            item.work_month or "",
            item.work_hours,
            item.overtime_hours,
            item.record_count or 0,
            item.approval_status or "",
        ]
    if dimension == "project":
        return [
            item.project_name or "",
            item.work_hours,
            item.overtime_hours,
            item.record_count or 0,
            item.approval_status or "",
        ]
    return [
        item.department,
        item.work_hours,
        item.overtime_hours,
        item.record_count or 0,
        item.approval_status or "",
    ]


@router.get(
    "/employee-work-hours",
    response_model=EmployeeWorkHourReportListResponse,
    summary="员工工时",
    description=(
        "查询员工工时数据，支持日期范围、部门、员工、项目筛选。"
        "统计维度：detail（明细）、employee（按员工汇总）、employee_date（按员工+日期）、"
        "employee_month（按员工+月份）、project（按项目）、department（按部门）。"
    ),
)
def list_employee_work_hours_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    date_from: date | None = Query(None, description="日期起（含）"),
    date_to: date | None = Query(None, description="日期止（含）"),
    department: str | None = Query(None, description="部门筛选"),
    employee_no: str | None = Query(None, description="工号筛选"),
    project_name: str | None = Query(None, description="项目筛选"),
    dimension: str = Query("detail", description="统计维度"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if dimension not in WORK_HOUR_DIMENSIONS:
        dimension = "detail"
    date_from, date_to = _normalize_work_hour_date_range(date_from, date_to)

    all_items = _build_work_hour_report_items(
        db,
        date_from=date_from,
        date_to=date_to,
        department=department,
        employee_no=employee_no,
        project_name=project_name,
        dimension=dimension,
    )
    total = len(all_items)
    work_hours_sum = round(sum(i.work_hours for i in all_items), 2)
    overtime_hours_sum = round(sum(i.overtime_hours for i in all_items), 2)
    start = (page - 1) * page_size
    page_items = all_items[start : start + page_size]

    return EmployeeWorkHourReportListResponse(
        items=page_items,
        total=total,
        page=page,
        page_size=page_size,
        dimension=dimension,
        work_hours_sum=work_hours_sum,
        overtime_hours_sum=overtime_hours_sum,
    )


@router.get(
    "/employee-work-hours/filters",
    response_model=EmployeeWorkHourFiltersResponse,
    summary="员工工时报表筛选选项",
    description="返回部门、员工、项目下拉选项。",
)
def list_employee_work_hour_filters(
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    departments = [
        name
        for (name,) in db.query(EmployeeWorkHour.department)
        .distinct()
        .order_by(EmployeeWorkHour.department)
        .all()
    ]
    employee_rows = (
        db.query(EmployeeWorkHour.employee_no, EmployeeWorkHour.employee_name)
        .distinct()
        .order_by(EmployeeWorkHour.employee_no)
        .all()
    )
    employees = [
        EmployeeWorkHourFilterEmployee(employee_no=no, employee_name=name)
        for no, name in employee_rows
    ]
    projects = [
        name
        for (name,) in db.query(EmployeeWorkHour.project_name)
        .distinct()
        .order_by(EmployeeWorkHour.project_name)
        .all()
    ]
    return EmployeeWorkHourFiltersResponse(
        departments=departments,
        employees=employees,
        projects=projects,
    )


@router.get(
    "/employee-work-hours/export",
    summary="导出员工工时报表 Excel",
    description="按当前筛选条件与统计维度导出 Excel 文件。",
)
def export_employee_work_hours_report(
    date_from: date | None = Query(None, description="日期起（含）"),
    date_to: date | None = Query(None, description="日期止（含）"),
    department: str | None = Query(None, description="部门筛选"),
    employee_no: str | None = Query(None, description="工号筛选"),
    project_name: str | None = Query(None, description="项目筛选"),
    dimension: str = Query("detail", description="统计维度"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if dimension not in WORK_HOUR_DIMENSIONS:
        dimension = "detail"
    date_from, date_to = _normalize_work_hour_date_range(date_from, date_to)

    items = _build_work_hour_report_items(
        db,
        date_from=date_from,
        date_to=date_to,
        department=department,
        employee_no=employee_no,
        project_name=project_name,
        dimension=dimension,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "员工工时报表"
    headers = _work_hour_export_headers(dimension)
    ws.append(headers)
    for item in items:
        ws.append(_work_hour_export_row(item, dimension))

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"employee_work_hours_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


REPAIR_STATUS_LABELS = {
    "pending": "待处理",
    "in_progress": "处理中",
    "completed": "已完成",
    "closed": "已关闭",
}

REPAIR_URGENCY_LABELS = {
    "low": "低",
    "normal": "普通",
    "high": "高",
    "urgent": "紧急",
}


def _repair_status_label(status: str | None) -> str:
    if not status:
        return "—"
    return REPAIR_STATUS_LABELS.get(status, status)


def _repair_urgency_label(urgency: str | None) -> str:
    if not urgency:
        return "—"
    return REPAIR_URGENCY_LABELS.get(urgency, urgency)


def _repair_duration_minutes(repair: EquipmentRepair) -> float | None:
    start = repair.start_time or repair.created_at
    if not start:
        return None
    end = repair.repair_completed_at
    if not end:
        if repair.status in ("in_progress", "pending") and repair.start_time:
            end = datetime.utcnow()
        else:
            return None
    minutes = (end - start).total_seconds() / 60
    if minutes < 0:
        return None
    return round(minutes, 1)


def _format_repair_duration(minutes: float | None) -> str:
    if minutes is None:
        return "—"
    total = int(round(minutes))
    if total < 60:
        return f"{total}分钟"
    hours, mins = divmod(total, 60)
    if mins == 0:
        return f"{hours}小时"
    return f"{hours}小时{mins}分钟"


def _repair_to_list_item(repair: EquipmentRepair) -> EquipmentRepairListItem:
    equipment = repair.equipment
    return EquipmentRepairListItem(
        id=repair.id,
        repair_no=repair.repair_no,
        equipment_id=repair.equipment_id,
        equipment_code=equipment.equipment_code if equipment else None,
        equipment_name=equipment.name if equipment else None,
        fault_category=repair.fault_category,
        fault_description=repair.fault_description,
        urgency=repair.urgency,
        status=repair.status,
        reporter=repair.reporter,
        repair_person=repair.repair_person,
        fault_time=repair.created_at,
        repair_duration_minutes=_repair_duration_minutes(repair),
        repair_completed_at=repair.repair_completed_at,
        created_at=repair.created_at,
    )


def _repair_to_detail(repair: EquipmentRepair) -> EquipmentRepairDetail:
    equipment = repair.equipment
    parts = [
        EquipmentRepairPartResponse(
            id=p.id,
            repair_id=p.repair_id,
            part_name=p.part_name,
            part_spec=p.part_spec,
            quantity=p.quantity,
            unit=p.unit,
            unit_price=float(p.unit_price or 0),
        )
        for p in (repair.parts or [])
    ]
    return EquipmentRepairDetail(
        id=repair.id,
        repair_no=repair.repair_no,
        equipment_id=repair.equipment_id,
        equipment_code=equipment.equipment_code if equipment else None,
        equipment_name=equipment.name if equipment else None,
        fault_category=repair.fault_category,
        fault_description=repair.fault_description,
        urgency=repair.urgency,
        status=repair.status,
        reporter=repair.reporter,
        repair_person=repair.repair_person,
        start_time=repair.start_time,
        repair_completed_at=repair.repair_completed_at,
        repair_description=repair.repair_description,
        images=repair.images,
        parts=parts,
        created_at=repair.created_at,
        updated_at=repair.updated_at,
    )


def _apply_equipment_repair_report_filters(
    query,
    *,
    keyword: str | None,
    status: str | None,
    date_from: date | None,
    date_to: date | None,
    equipment_code: str | None,
    fault_category: str | None,
):
    needs_equipment_join = bool(keyword or equipment_code)
    if needs_equipment_join:
        query = query.join(Equipment)
    if keyword:
        pattern = f"%{keyword}%"
        query = query.filter(
            (EquipmentRepair.repair_no.ilike(pattern))
            | (Equipment.name.ilike(pattern))
            | (Equipment.equipment_code.ilike(pattern))
            | (EquipmentRepair.fault_description.ilike(pattern))
        )
    if equipment_code:
        query = query.filter(Equipment.equipment_code == equipment_code)
    if status:
        query = query.filter(EquipmentRepair.status == status)
    if fault_category:
        query = query.filter(EquipmentRepair.fault_category == fault_category)
    if date_from:
        day_start = datetime.combine(date_from, datetime.min.time())
        query = query.filter(EquipmentRepair.created_at >= day_start)
    if date_to:
        day_end = datetime.combine(date_to + timedelta(days=1), datetime.min.time())
        query = query.filter(EquipmentRepair.created_at < day_end)
    return query


def _parts_cost_total(repair: EquipmentRepair) -> float:
    return round(
        sum(float(p.unit_price or 0) * int(p.quantity or 0) for p in (repair.parts or [])),
        2,
    )


@router.get(
    "/equipment-repairs",
    response_model=EquipmentRepairListResponse,
    summary="设备维修报表",
    description=(
        "报表中心设备维修列表：分页查询维修工单，"
        "支持关键字、状态、报修日期范围、设备编号、故障分类筛选。"
    ),
)
def list_equipment_repair_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    keyword: str | None = Query(None, description="工单号/设备/故障描述关键字"),
    status: str | None = Query(None, description="工单状态"),
    date_from: date | None = Query(None, description="报修日期起（含）"),
    date_to: date | None = Query(None, description="报修日期止（含）"),
    equipment_code: str | None = Query(None, description="设备编号（精确匹配）"),
    fault_category: str | None = Query(None, description="故障分类"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(EquipmentRepair)
    query = _apply_equipment_repair_report_filters(
        query,
        keyword=keyword,
        status=status,
        date_from=date_from,
        date_to=date_to,
        equipment_code=equipment_code,
        fault_category=fault_category,
    )
    query = query.order_by(EquipmentRepair.created_at.desc(), EquipmentRepair.id.desc())
    total = query.count()
    rows = query.offset((page - 1) * page_size).limit(page_size).all()
    return EquipmentRepairListResponse(
        items=[_repair_to_list_item(r) for r in rows],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get(
    "/equipment-repairs/export",
    summary="导出设备维修报表 Excel",
    description="按当前筛选条件导出设备维修工单 Excel 文件。",
)
def export_equipment_repair_report(
    keyword: str | None = Query(None, description="工单号/设备/故障描述关键字"),
    status: str | None = Query(None, description="工单状态"),
    date_from: date | None = Query(None, description="报修日期起（含）"),
    date_to: date | None = Query(None, description="报修日期止（含）"),
    equipment_code: str | None = Query(None, description="设备编号（精确匹配）"),
    fault_category: str | None = Query(None, description="故障分类"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = db.query(EquipmentRepair)
    query = _apply_equipment_repair_report_filters(
        query,
        keyword=keyword,
        status=status,
        date_from=date_from,
        date_to=date_to,
        equipment_code=equipment_code,
        fault_category=fault_category,
    )
    rows = query.order_by(EquipmentRepair.created_at.desc(), EquipmentRepair.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "设备维修报表"
    headers = [
        "工单号",
        "设备编号",
        "设备名称",
        "故障时间",
        "故障现象",
        "维修人",
        "耗时",
        "状态",
        "故障分类",
        "报修人",
        "完成时间",
    ]
    ws.append(headers)
    for repair in rows:
        equipment = repair.equipment
        ws.append(
            [
                repair.repair_no,
                equipment.equipment_code if equipment else "",
                equipment.name if equipment else "",
                repair.created_at.strftime("%Y-%m-%d %H:%M:%S") if repair.created_at else "",
                repair.fault_description,
                repair.repair_person or "",
                _format_repair_duration(_repair_duration_minutes(repair)),
                _repair_status_label(repair.status),
                repair.fault_category,
                repair.reporter,
                repair.repair_completed_at.strftime("%Y-%m-%d %H:%M:%S")
                if repair.repair_completed_at
                else "",
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"equipment_repair_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/equipment-repairs/{repair_id}",
    response_model=EquipmentRepairDetail,
    summary="设备维修报表详情",
    description="按 ID 返回维修工单详情（含配件明细），供报表中心详情查看。",
)
def get_equipment_repair_report_detail(
    repair_id: int,
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    repair = (
        db.query(EquipmentRepair)
        .filter(EquipmentRepair.id == repair_id)
        .first()
    )
    if not repair:
        raise HTTPException(status_code=404, detail="维修工单不存在")
    return _repair_to_detail(repair)


def _normalize_equipment_oee_date_range(
    date_from: date | None, date_to: date | None
) -> tuple[date, date]:
    today = date.today()
    if date_to is None:
        date_to = today
    if date_from is None:
        date_from = date_to - timedelta(days=6)
    if date_from > date_to:
        date_from, date_to = date_to, date_from
    return date_from, date_to


def _empty_equipment_oee_summary() -> EquipmentOeeSummary:
    return EquipmentOeeSummary(
        oee=0,
        availability=0,
        performance=0,
        quality=0,
        utilization_rate=0,
        startup_rate=0,
        downtime_hours=0,
    )


def _filter_equipment_for_oee_report(
    db: Session,
    workshop: str | None,
    equipment_type: str | None,
    equipment_id: int | None,
) -> list[Equipment]:
    query = db.query(Equipment)
    if workshop:
        query = query.filter(Equipment.department == workshop)
    if equipment_type:
        query = query.filter(Equipment.spec_model == equipment_type)
    if equipment_id:
        query = query.filter(Equipment.id == equipment_id)
    return query.order_by(Equipment.equipment_code).all()


def _build_equipment_oee_report_items(
    db: Session,
    date_from: date,
    date_to: date,
    workshop: str | None,
    equipment_type: str | None,
    equipment_id: int | None,
) -> list[EquipmentOeeReportItem]:
    equipment_list = _filter_equipment_for_oee_report(
        db, workshop=workshop, equipment_type=equipment_type, equipment_id=equipment_id
    )
    if not equipment_list:
        return []

    equipment_by_id = {eq.id: eq for eq in equipment_list}
    equipment_ids = list(equipment_by_id.keys())

    oee_rows = (
        db.query(EquipmentOeeSnapshot)
        .filter(
            EquipmentOeeSnapshot.equipment_id.in_(equipment_ids),
            EquipmentOeeSnapshot.period_type == "day",
            EquipmentOeeSnapshot.period_start >= date_from,
            EquipmentOeeSnapshot.period_start <= date_to,
        )
        .all()
    )
    oee_map: dict[tuple[int, date], EquipmentOeeSnapshot] = {
        (row.equipment_id, row.period_start): row for row in oee_rows
    }

    runtime_rows = (
        db.query(EquipmentRuntimeLog)
        .filter(
            EquipmentRuntimeLog.equipment_id.in_(equipment_ids),
            func.date(EquipmentRuntimeLog.start_at) >= date_from,
            func.date(EquipmentRuntimeLog.start_at) <= date_to,
        )
        .all()
    )
    runtime_stats: dict[tuple[int, date], dict[str, float]] = defaultdict(
        lambda: {"running": 0.0, "stop": 0.0, "total": 0.0}
    )
    for row in runtime_rows:
        day = row.start_at.date()
        key = (row.equipment_id, day)
        hours = float(row.runtime_hours or 0)
        runtime_stats[key]["total"] += hours
        if row.status == "运行":
            runtime_stats[key]["running"] += hours
        elif row.status == "停机":
            runtime_stats[key]["stop"] += hours

    output_rows = (
        db.query(EquipmentOutputRecord)
        .filter(
            EquipmentOutputRecord.equipment_id.in_(equipment_ids),
            EquipmentOutputRecord.record_date >= date_from,
            EquipmentOutputRecord.record_date <= date_to,
        )
        .all()
    )
    output_map: dict[tuple[int, date], int] = {
        (row.equipment_id, row.record_date): int(row.output_qty or 0) for row in output_rows
    }

    items: list[EquipmentOeeReportItem] = []
    current = date_from
    while current <= date_to:
        for eq_id, eq in equipment_by_id.items():
            snap = oee_map.get((eq_id, current))
            runtime = runtime_stats.get((eq_id, current), {"running": 0.0, "stop": 0.0, "total": 0.0})
            total_hours = runtime["total"]
            running_hours = runtime["running"]
            stop_hours = runtime["stop"]
            utilization_rate = round(running_hours / total_hours * 100, 2) if total_hours else 0.0
            startup_rate = 100.0 if running_hours > 0 else 0.0

            if snap:
                availability = float(snap.availability)
                performance = float(snap.performance)
                quality = float(snap.quality)
                oee = float(snap.oee)
            else:
                availability = performance = quality = oee = 0.0

            has_data = bool(snap) or total_hours > 0 or (eq_id, current) in output_map
            if not has_data:
                continue

            items.append(
                EquipmentOeeReportItem(
                    equipment_code=eq.equipment_code,
                    equipment_name=eq.name,
                    workshop=eq.department,
                    equipment_type=eq.spec_model,
                    period_date=current,
                    oee=oee,
                    availability=availability,
                    performance=performance,
                    quality=quality,
                    utilization_rate=utilization_rate,
                    startup_rate=startup_rate,
                    downtime_hours=round(stop_hours, 2),
                    output_qty=output_map.get((eq_id, current)),
                )
            )
        current += timedelta(days=1)

    items.sort(key=lambda item: (item.period_date, item.equipment_code), reverse=True)
    return items


def _average_metric(values: list[float]) -> float:
    if not values:
        return 0.0
    return round(sum(values) / len(values), 2)


def _build_equipment_oee_summary(items: list[EquipmentOeeReportItem]) -> EquipmentOeeSummary:
    if not items:
        return _empty_equipment_oee_summary()
    return EquipmentOeeSummary(
        oee=_average_metric([item.oee for item in items]),
        availability=_average_metric([item.availability for item in items]),
        performance=_average_metric([item.performance for item in items]),
        quality=_average_metric([item.quality for item in items]),
        utilization_rate=_average_metric([item.utilization_rate for item in items]),
        startup_rate=_average_metric([item.startup_rate for item in items]),
        downtime_hours=round(sum(item.downtime_hours for item in items), 2),
    )


def _build_equipment_oee_trend(items: list[EquipmentOeeReportItem]) -> list[EquipmentOeeTrendPoint]:
    grouped: dict[date, list[EquipmentOeeReportItem]] = defaultdict(list)
    for item in items:
        grouped[item.period_date].append(item)

    trend: list[EquipmentOeeTrendPoint] = []
    for period_date in sorted(grouped.keys()):
        rows = grouped[period_date]
        trend.append(
            EquipmentOeeTrendPoint(
                period_date=period_date,
                oee=_average_metric([row.oee for row in rows]),
                availability=_average_metric([row.availability for row in rows]),
                performance=_average_metric([row.performance for row in rows]),
                quality=_average_metric([row.quality for row in rows]),
                utilization_rate=_average_metric([row.utilization_rate for row in rows]),
                startup_rate=_average_metric([row.startup_rate for row in rows]),
                downtime_hours=round(sum(row.downtime_hours for row in rows), 2),
            )
        )
    return trend


@router.get(
    "/equipment-oee/filters",
    response_model=EquipmentOeeFiltersResponse,
    summary="设备 OEE 报表筛选选项",
    description="返回车间、设备类型、设备下拉选项。",
)
def list_equipment_oee_filters(
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    workshops = [
        name
        for (name,) in db.query(Equipment.department)
        .filter(Equipment.department.isnot(None), Equipment.department != "")
        .distinct()
        .order_by(Equipment.department)
        .all()
    ]
    equipment_types = [
        name
        for (name,) in db.query(Equipment.spec_model)
        .filter(Equipment.spec_model.isnot(None), Equipment.spec_model != "")
        .distinct()
        .order_by(Equipment.spec_model)
        .all()
    ]
    equipment_rows = db.query(Equipment).order_by(Equipment.equipment_code).all()
    equipment = [
        EquipmentOeeFilterEquipment(
            id=eq.id,
            equipment_code=eq.equipment_code,
            name=eq.name,
            workshop=eq.department,
            equipment_type=eq.spec_model,
        )
        for eq in equipment_rows
    ]
    return EquipmentOeeFiltersResponse(
        workshops=workshops,
        equipment_types=equipment_types,
        equipment=equipment,
    )


@router.get(
    "/equipment-oee",
    response_model=EquipmentOeeReportListResponse,
    summary="设备 OEE 报表",
    description=(
        "按时间范围、车间、设备类型/单台设备查询 OEE 与稼动指标，"
        "返回汇总指标、趋势序列与明细分页。"
    ),
)
def list_equipment_oee_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    date_from: date | None = Query(None, description="日期起（含）"),
    date_to: date | None = Query(None, description="日期止（含）"),
    workshop: str | None = Query(None, description="车间筛选"),
    equipment_type: str | None = Query(None, description="设备类型（型号）筛选"),
    equipment_id: int | None = Query(None, description="单台设备 ID"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    date_from, date_to = _normalize_equipment_oee_date_range(date_from, date_to)
    all_items = _build_equipment_oee_report_items(
        db,
        date_from=date_from,
        date_to=date_to,
        workshop=workshop,
        equipment_type=equipment_type,
        equipment_id=equipment_id,
    )
    total = len(all_items)
    summary = _build_equipment_oee_summary(all_items)
    trend = _build_equipment_oee_trend(all_items)
    start = (page - 1) * page_size
    page_items = all_items[start : start + page_size]
    return EquipmentOeeReportListResponse(
        items=page_items,
        total=total,
        page=page,
        page_size=page_size,
        summary=summary,
        trend=trend,
    )


@router.get(
    "/equipment-oee/export",
    summary="导出设备 OEE 报表 Excel",
    description="按当前筛选条件导出设备 OEE 明细 Excel 文件。",
)
def export_equipment_oee_report(
    date_from: date | None = Query(None, description="日期起（含）"),
    date_to: date | None = Query(None, description="日期止（含）"),
    workshop: str | None = Query(None, description="车间筛选"),
    equipment_type: str | None = Query(None, description="设备类型（型号）筛选"),
    equipment_id: int | None = Query(None, description="单台设备 ID"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    date_from, date_to = _normalize_equipment_oee_date_range(date_from, date_to)
    rows = _build_equipment_oee_report_items(
        db,
        date_from=date_from,
        date_to=date_to,
        workshop=workshop,
        equipment_type=equipment_type,
        equipment_id=equipment_id,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "设备OEE报表"
    headers = [
        "设备编号",
        "设备名称",
        "车间",
        "设备型号",
        "日期",
        "OEE(%)",
        "时间稼动率(%)",
        "性能稼动率(%)",
        "良率(%)",
        "稼动率(%)",
        "开机率(%)",
        "停机时长(h)",
        "产量",
    ]
    ws.append(headers)
    for item in rows:
        ws.append(
            [
                item.equipment_code,
                item.equipment_name,
                item.workshop or "",
                item.equipment_type or "",
                item.period_date.isoformat(),
                item.oee,
                item.availability,
                item.performance,
                item.quality,
                item.utilization_rate,
                item.startup_rate,
                item.downtime_hours,
                item.output_qty if item.output_qty is not None else "",
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"equipment_oee_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


DOWNTIME_STATUSES = ("停机", "维修", "待机")


def _normalize_equipment_downtime_date_range(
    date_from: date | None, date_to: date | None
) -> tuple[date, date]:
    today = date.today()
    if date_to is None:
        date_to = today
    if date_from is None:
        date_from = date_to - timedelta(days=6)
    if date_from > date_to:
        date_from, date_to = date_to, date_from
    return date_from, date_to


def _filter_equipment_for_downtime_report(
    db: Session,
    workshop: str | None,
    equipment_type: str | None,
    equipment_id: int | None,
) -> list[Equipment]:
    query = db.query(Equipment)
    if workshop:
        query = query.filter(Equipment.department == workshop)
    if equipment_type:
        query = query.filter(Equipment.spec_model == equipment_type)
    if equipment_id:
        query = query.filter(Equipment.id == equipment_id)
    return query.order_by(Equipment.equipment_code).all()


def _build_equipment_downtime_report_items(
    db: Session,
    date_from: date,
    date_to: date,
    workshop: str | None,
    equipment_type: str | None,
    equipment_id: int | None,
    status: str | None,
) -> list[EquipmentDowntimeReportItem]:
    equipment_list = _filter_equipment_for_downtime_report(
        db, workshop=workshop, equipment_type=equipment_type, equipment_id=equipment_id
    )
    if not equipment_list:
        return []

    equipment_by_id = {eq.id: eq for eq in equipment_list}
    equipment_ids = list(equipment_by_id.keys())

    query = (
        db.query(EquipmentRuntimeLog)
        .filter(
            EquipmentRuntimeLog.equipment_id.in_(equipment_ids),
            EquipmentRuntimeLog.status.in_(DOWNTIME_STATUSES),
            func.date(EquipmentRuntimeLog.start_at) >= date_from,
            func.date(EquipmentRuntimeLog.start_at) <= date_to,
        )
        .order_by(EquipmentRuntimeLog.start_at.desc(), EquipmentRuntimeLog.id.desc())
    )
    if status:
        query = query.filter(EquipmentRuntimeLog.status == status)

    items: list[EquipmentDowntimeReportItem] = []
    for row in query.all():
        eq = equipment_by_id.get(row.equipment_id)
        if not eq:
            continue
        items.append(
            EquipmentDowntimeReportItem(
                id=row.id,
                equipment_id=row.equipment_id,
                equipment_code=eq.equipment_code,
                equipment_name=eq.name,
                workshop=eq.department,
                equipment_type=eq.spec_model,
                start_at=row.start_at,
                end_at=row.end_at,
                status=row.status,
                downtime_hours=round(float(row.runtime_hours or 0), 2),
            )
        )
    return items


def _empty_equipment_downtime_summary() -> EquipmentDowntimeSummary:
    return EquipmentDowntimeSummary(
        event_count=0,
        total_downtime_hours=0,
        avg_downtime_hours=0,
        equipment_count=0,
    )


def _build_equipment_downtime_summary(
    items: list[EquipmentDowntimeReportItem],
) -> EquipmentDowntimeSummary:
    if not items:
        return _empty_equipment_downtime_summary()
    total_hours = round(sum(item.downtime_hours for item in items), 2)
    equipment_count = len({item.equipment_id for item in items})
    event_count = len(items)
    avg_hours = round(total_hours / event_count, 2) if event_count else 0.0
    return EquipmentDowntimeSummary(
        event_count=event_count,
        total_downtime_hours=total_hours,
        avg_downtime_hours=avg_hours,
        equipment_count=equipment_count,
    )


def _build_equipment_downtime_trend(
    items: list[EquipmentDowntimeReportItem],
) -> list[EquipmentDowntimeTrendPoint]:
    grouped: dict[date, list[EquipmentDowntimeReportItem]] = defaultdict(list)
    for item in items:
        grouped[item.start_at.date()].append(item)

    trend: list[EquipmentDowntimeTrendPoint] = []
    for period_date in sorted(grouped.keys()):
        rows = grouped[period_date]
        trend.append(
            EquipmentDowntimeTrendPoint(
                period_date=period_date,
                event_count=len(rows),
                downtime_hours=round(sum(row.downtime_hours for row in rows), 2),
            )
        )
    return trend


def _shift_label(start_at: datetime) -> str:
    hour = start_at.hour
    if 8 <= hour < 16:
        return "早班"
    if 16 <= hour < 24:
        return "中班"
    return "夜班"


def _period_calendar_hours(date_from: date, date_to: date) -> float:
    return float((date_to - date_from).days + 1) * 24.0


def _build_equipment_downtime_dimension_stats(
    items: list[EquipmentDowntimeReportItem],
    key_fn,
    label_fn,
) -> list[EquipmentDowntimeDimensionStat]:
    grouped: dict[str, list[EquipmentDowntimeReportItem]] = defaultdict(list)
    labels: dict[str, str] = {}
    for item in items:
        key = key_fn(item)
        grouped[key].append(item)
        labels[key] = label_fn(item)

    total_hours = round(sum(item.downtime_hours for item in items), 2)
    stats: list[EquipmentDowntimeDimensionStat] = []
    for key, rows in grouped.items():
        hours = round(sum(row.downtime_hours for row in rows), 2)
        stats.append(
            EquipmentDowntimeDimensionStat(
                dimension_key=key,
                dimension_label=labels[key],
                event_count=len(rows),
                downtime_hours=hours,
                downtime_pct=round(hours / total_hours * 100, 2) if total_hours else 0.0,
            )
        )
    return sorted(stats, key=lambda row: row.downtime_hours, reverse=True)


def _build_equipment_downtime_by_equipment(
    items: list[EquipmentDowntimeReportItem],
) -> list[EquipmentDowntimeDimensionStat]:
    return _build_equipment_downtime_dimension_stats(
        items,
        key_fn=lambda item: str(item.equipment_id),
        label_fn=lambda item: f"{item.equipment_name}（{item.equipment_code}）",
    )


def _build_equipment_downtime_by_line(
    items: list[EquipmentDowntimeReportItem],
) -> list[EquipmentDowntimeDimensionStat]:
    return _build_equipment_downtime_dimension_stats(
        items,
        key_fn=lambda item: item.workshop or "未分配",
        label_fn=lambda item: item.workshop or "未分配",
    )


def _build_equipment_downtime_by_shift(
    items: list[EquipmentDowntimeReportItem],
) -> list[EquipmentDowntimeDimensionStat]:
    return _build_equipment_downtime_dimension_stats(
        items,
        key_fn=lambda item: _shift_label(item.start_at),
        label_fn=lambda item: _shift_label(item.start_at),
    )


def _build_equipment_downtime_reason_pareto(
    items: list[EquipmentDowntimeReportItem],
) -> list[EquipmentDowntimeReasonParetoItem]:
    grouped: dict[str, list[EquipmentDowntimeReportItem]] = defaultdict(list)
    for item in items:
        grouped[item.status].append(item)

    total_hours = round(sum(item.downtime_hours for item in items), 2)
    ranked = sorted(
        grouped.items(),
        key=lambda pair: sum(row.downtime_hours for row in pair[1]),
        reverse=True,
    )
    pareto: list[EquipmentDowntimeReasonParetoItem] = []
    cumulative = 0.0
    for reason, rows in ranked:
        hours = round(sum(row.downtime_hours for row in rows), 2)
        cumulative += hours
        pareto.append(
            EquipmentDowntimeReasonParetoItem(
                reason=reason,
                event_count=len(rows),
                downtime_hours=hours,
                cumulative_pct=round(cumulative / total_hours * 100, 2) if total_hours else 0.0,
            )
        )
    return pareto


def _build_equipment_downtime_reliability(
    items: list[EquipmentDowntimeReportItem],
    date_from: date,
    date_to: date,
) -> list[EquipmentDowntimeReliabilityMetric]:
    period_hours = _period_calendar_hours(date_from, date_to)
    grouped: dict[int, list[EquipmentDowntimeReportItem]] = defaultdict(list)
    for item in items:
        grouped[item.equipment_id].append(item)

    metrics: list[EquipmentDowntimeReliabilityMetric] = []
    for equipment_id, rows in grouped.items():
        total_downtime = round(sum(row.downtime_hours for row in rows), 2)
        event_count = len(rows)
        mttr = round(total_downtime / event_count, 2) if event_count else 0.0
        operating_hours = max(period_hours - total_downtime, 0.0)
        mtbf = round(operating_hours / event_count, 2) if event_count else 0.0
        sample = rows[0]
        metrics.append(
            EquipmentDowntimeReliabilityMetric(
                equipment_id=equipment_id,
                equipment_code=sample.equipment_code,
                equipment_name=sample.equipment_name,
                event_count=event_count,
                mtbf_hours=mtbf,
                mttr_hours=mttr,
            )
        )
    return sorted(metrics, key=lambda row: row.event_count, reverse=True)


def _append_dimension_sheet(ws, title: str, rows: list[EquipmentDowntimeDimensionStat]) -> None:
    ws.title = title
    ws.append(["名称", "停机次数", "停机时长(h)", "时长占比(%)"])
    for row in rows:
        ws.append([row.dimension_label, row.event_count, row.downtime_hours, row.downtime_pct])


def _append_pareto_sheet(ws, rows: list[EquipmentDowntimeReasonParetoItem]) -> None:
    ws.title = "原因Pareto"
    ws.append(["停机原因", "停机次数", "停机时长(h)", "累计占比(%)"])
    for row in rows:
        ws.append([row.reason, row.event_count, row.downtime_hours, row.cumulative_pct])


def _append_reliability_sheet(ws, rows: list[EquipmentDowntimeReliabilityMetric]) -> None:
    ws.title = "MTBF_MTTR"
    ws.append(["设备编号", "设备名称", "停机次数", "MTBF(h)", "MTTR(h)"])
    for row in rows:
        ws.append(
            [
                row.equipment_code,
                row.equipment_name,
                row.event_count,
                row.mtbf_hours,
                row.mttr_hours,
            ]
        )


@router.get(
    "/equipment-downtime/filters",
    response_model=EquipmentDowntimeFiltersResponse,
    summary="设备停机报表筛选选项",
    description="返回车间、设备类型、设备列表与停机类型下拉选项。",
)
def list_equipment_downtime_filters(
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    workshops = [
        name
        for (name,) in db.query(Equipment.department)
        .filter(Equipment.department.isnot(None), Equipment.department != "")
        .distinct()
        .order_by(Equipment.department)
        .all()
    ]
    equipment_types = [
        name
        for (name,) in db.query(Equipment.spec_model)
        .filter(Equipment.spec_model.isnot(None), Equipment.spec_model != "")
        .distinct()
        .order_by(Equipment.spec_model)
        .all()
    ]
    equipment_rows = db.query(Equipment).order_by(Equipment.equipment_code).all()
    equipment = [
        EquipmentDowntimeFilterEquipment(
            id=eq.id,
            equipment_code=eq.equipment_code,
            name=eq.name,
            workshop=eq.department,
            equipment_type=eq.spec_model,
        )
        for eq in equipment_rows
    ]
    return EquipmentDowntimeFiltersResponse(
        workshops=workshops,
        equipment_types=equipment_types,
        equipment=equipment,
        downtime_statuses=list(DOWNTIME_STATUSES),
    )


@router.get(
    "/equipment-downtime",
    response_model=EquipmentDowntimeReportListResponse,
    summary="设备停机报表",
    description=(
        "按时间范围、车间、设备类型/单台设备、停机类型查询设备停机明细，"
        "数据来自 equipment_runtime_logs（状态为停机/维修/待机）。"
        "返回汇总指标、趋势序列与明细分页。"
    ),
)
def list_equipment_downtime_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    date_from: date | None = Query(None, description="日期起（含）"),
    date_to: date | None = Query(None, description="日期止（含）"),
    workshop: str | None = Query(None, description="车间筛选"),
    equipment_type: str | None = Query(None, description="设备类型（型号）筛选"),
    equipment_id: int | None = Query(None, description="单台设备 ID"),
    status: str | None = Query(None, description="停机类型：停机/维修/待机"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if status and status not in DOWNTIME_STATUSES:
        status = None
    date_from, date_to = _normalize_equipment_downtime_date_range(date_from, date_to)
    all_items = _build_equipment_downtime_report_items(
        db,
        date_from=date_from,
        date_to=date_to,
        workshop=workshop,
        equipment_type=equipment_type,
        equipment_id=equipment_id,
        status=status,
    )
    total = len(all_items)
    summary = _build_equipment_downtime_summary(all_items)
    trend = _build_equipment_downtime_trend(all_items)
    by_equipment = _build_equipment_downtime_by_equipment(all_items)
    by_line = _build_equipment_downtime_by_line(all_items)
    by_shift = _build_equipment_downtime_by_shift(all_items)
    reason_pareto = _build_equipment_downtime_reason_pareto(all_items)
    reliability = _build_equipment_downtime_reliability(all_items, date_from, date_to)
    start = (page - 1) * page_size
    page_items = all_items[start : start + page_size]
    return EquipmentDowntimeReportListResponse(
        items=page_items,
        total=total,
        page=page,
        page_size=page_size,
        summary=summary,
        trend=trend,
        by_equipment=by_equipment,
        by_line=by_line,
        by_shift=by_shift,
        reason_pareto=reason_pareto,
        reliability=reliability,
    )


@router.get(
    "/equipment-downtime/export",
    summary="导出设备停机报表 Excel",
    description="按当前筛选条件导出设备停机明细 Excel 文件。",
)
def export_equipment_downtime_report(
    date_from: date | None = Query(None, description="日期起（含）"),
    date_to: date | None = Query(None, description="日期止（含）"),
    workshop: str | None = Query(None, description="车间筛选"),
    equipment_type: str | None = Query(None, description="设备类型（型号）筛选"),
    equipment_id: int | None = Query(None, description="单台设备 ID"),
    status: str | None = Query(None, description="停机类型：停机/维修/待机"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    if status and status not in DOWNTIME_STATUSES:
        status = None
    date_from, date_to = _normalize_equipment_downtime_date_range(date_from, date_to)
    rows = _build_equipment_downtime_report_items(
        db,
        date_from=date_from,
        date_to=date_to,
        workshop=workshop,
        equipment_type=equipment_type,
        equipment_id=equipment_id,
        status=status,
    )
    by_equipment = _build_equipment_downtime_by_equipment(rows)
    by_line = _build_equipment_downtime_by_line(rows)
    by_shift = _build_equipment_downtime_by_shift(rows)
    reason_pareto = _build_equipment_downtime_reason_pareto(rows)
    reliability = _build_equipment_downtime_reliability(rows, date_from, date_to)

    wb = Workbook()
    ws_detail = wb.active
    ws_detail.title = "停机明细"
    headers = [
        "设备编号",
        "设备名称",
        "车间",
        "设备型号",
        "停机类型",
        "开始时间",
        "结束时间",
        "停机时长(h)",
    ]
    ws_detail.append(headers)
    for item in rows:
        ws_detail.append(
            [
                item.equipment_code,
                item.equipment_name,
                item.workshop or "",
                item.equipment_type or "",
                item.status,
                item.start_at.strftime("%Y-%m-%d %H:%M:%S") if item.start_at else "",
                item.end_at.strftime("%Y-%m-%d %H:%M:%S") if item.end_at else "",
                item.downtime_hours,
            ]
        )

    _append_dimension_sheet(wb.create_sheet(), "按设备统计", by_equipment)
    _append_dimension_sheet(wb.create_sheet(), "按产线统计", by_line)
    _append_dimension_sheet(wb.create_sheet(), "按班次统计", by_shift)
    _append_pareto_sheet(wb.create_sheet(), reason_pareto)
    _append_reliability_sheet(wb.create_sheet(), reliability)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"downtime_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _inspection_item_to_report_row(
    item: InspectionRecordItem,
    record: InspectionRecord,
    device: Device,
) -> EquipmentInspectionReportItem:
    return EquipmentInspectionReportItem(
        id=item.id,
        record_id=record.id,
        device_id=device.id,
        device_code=device.code,
        device_name=device.name,
        workshop=device.location,
        item_name=item.item_name,
        standard_value=item.standard_value,
        actual_value=item.actual_value,
        result=item.result,
        inspector=record.inspector,
        inspect_date=record.inspect_date,
        record_status=record.status,
        item_remark=item.remark,
    )


def _apply_equipment_inspection_report_filters(
    query,
    *,
    date_from: date | None,
    date_to: date | None,
    device_id: int | None,
    workshop: str | None,
    status: str | None,
):
    if date_from:
        query = query.filter(InspectionRecord.inspect_date >= date_from)
    if date_to:
        query = query.filter(InspectionRecord.inspect_date <= date_to)
    if device_id:
        query = query.filter(InspectionRecord.device_id == device_id)
    if workshop:
        query = query.filter(Device.location == workshop)
    if status:
        query = query.filter(InspectionRecord.status == status)
    return query


@router.get(
    "/equipment-inspection/filters",
    response_model=EquipmentInspectionFiltersResponse,
    summary="设备点检报表筛选选项",
    description="返回车间（设备位置）与点检设备下拉选项。",
)
def list_equipment_inspection_filters(
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    workshops = [
        loc
        for (loc,) in db.query(Device.location)
        .filter(Device.location.isnot(None), Device.location != "")
        .distinct()
        .order_by(Device.location)
        .all()
    ]
    device_rows = (
        db.query(Device.id, Device.code, Device.name, Device.location)
        .order_by(Device.code)
        .all()
    )
    devices = [
        EquipmentInspectionFilterDevice(
            id=row_id,
            code=code,
            name=name,
            workshop=location,
        )
        for row_id, code, name, location in device_rows
    ]
    return EquipmentInspectionFiltersResponse(workshops=workshops, devices=devices)


@router.get(
    "/equipment-inspection",
    response_model=EquipmentInspectionReportListResponse,
    summary="设备点检报表",
    description=(
        "点检明细报表：按点检项展开，支持日期区间、设备、车间（设备位置）、"
        "点检状态筛选与分页。"
    ),
)
def list_equipment_inspection_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    date_from: date | None = Query(None, description="点检日期起（含）"),
    date_to: date | None = Query(None, description="点检日期止（含）"),
    device_id: int | None = Query(None, description="点检设备 ID"),
    workshop: str | None = Query(None, description="车间（设备位置）"),
    status: str | None = Query(None, description="点检记录状态"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = (
        db.query(InspectionRecordItem, InspectionRecord, Device)
        .join(InspectionRecord, InspectionRecordItem.record_id == InspectionRecord.id)
        .join(Device, InspectionRecord.device_id == Device.id)
    )
    query = _apply_equipment_inspection_report_filters(
        query,
        date_from=date_from,
        date_to=date_to,
        device_id=device_id,
        workshop=workshop,
        status=status,
    )
    query = query.order_by(
        InspectionRecord.inspect_date.desc(),
        InspectionRecord.id.desc(),
        InspectionRecordItem.id.asc(),
    )
    total = query.count()
    rows = query.offset((page - 1) * page_size).limit(page_size).all()
    items = [
        _inspection_item_to_report_row(item, record, device)
        for item, record, device in rows
    ]
    return EquipmentInspectionReportListResponse(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
    )


def _maintenance_order_status_label(status: str) -> str:
    labels = {
        "pending": "待保养",
        "in_progress": "保养中",
        "completed": "已完成",
        "closed": "已关闭",
    }
    return labels.get(status, status)


def _apply_equipment_maintenance_report_filters(
    query,
    *,
    keyword: str | None,
    status: str | None,
    date_from: date | None,
    date_to: date | None,
    equipment_code: str | None,
):
    needs_equipment_join = bool(keyword or equipment_code)
    if needs_equipment_join:
        query = query.join(Equipment)
    if keyword:
        pattern = f"%{keyword}%"
        query = query.filter(
            (EquipmentMaintenanceOrder.order_no.ilike(pattern))
            | (Equipment.name.ilike(pattern))
            | (Equipment.equipment_code.ilike(pattern))
            | (EquipmentMaintenanceOrder.assignee.ilike(pattern))
            | (EquipmentMaintenanceOrder.executor.ilike(pattern))
        )
    if status:
        query = query.filter(EquipmentMaintenanceOrder.status == status)
    if equipment_code:
        query = query.filter(Equipment.equipment_code == equipment_code)
    if date_from:
        day_start = datetime.combine(date_from, datetime.min.time())
        query = query.filter(EquipmentMaintenanceOrder.planned_start_at >= day_start)
    if date_to:
        day_end = datetime.combine(date_to + timedelta(days=1), datetime.min.time())
        query = query.filter(EquipmentMaintenanceOrder.planned_start_at < day_end)
    return query


def _calc_maintenance_report_summary(
    orders: list[EquipmentMaintenanceOrder],
) -> EquipmentMaintenanceReportSummary:
    due_total = len(orders)
    completed = sum(1 for order in orders if order.status == "completed")
    not_done = sum(1 for order in orders if order.status in ("pending", "in_progress"))
    completion_rate = round(completed / due_total * 100, 1) if due_total else 0.0
    return EquipmentMaintenanceReportSummary(
        due_total=due_total,
        completed=completed,
        not_done=not_done,
        completion_rate=completion_rate,
    )


def _flatten_maintenance_orders_for_report(
    orders: list[EquipmentMaintenanceOrder],
) -> list[EquipmentMaintenanceReportItem]:
    rows: list[EquipmentMaintenanceReportItem] = []
    row_id = 1
    sorted_orders = sorted(
        orders,
        key=lambda order: (order.planned_start_at, order.id),
        reverse=True,
    )
    for order in sorted_orders:
        equipment = order.equipment
        plan = order.plan
        maintainer = order.executor or order.assignee
        maintenance_time = (
            order.actual_end_at or order.actual_start_at or order.planned_start_at
        )
        results = order.results or []
        if results:
            for item in results:
                if isinstance(item, dict):
                    item_name = item.get("item_name")
                    item_result = item.get("result")
                else:
                    item_name = getattr(item, "item_name", None)
                    item_result = getattr(item, "result", None)
                rows.append(
                    EquipmentMaintenanceReportItem(
                        id=row_id,
                        order_id=order.id,
                        order_no=order.order_no,
                        equipment_code=equipment.equipment_code if equipment else None,
                        equipment_name=equipment.name if equipment else None,
                        plan_name=plan.name if plan else None,
                        maintainer=maintainer,
                        maintenance_time=maintenance_time,
                        item_name=item_name,
                        item_result=item_result,
                        order_status=order.status,
                    )
                )
                row_id += 1
        else:
            rows.append(
                EquipmentMaintenanceReportItem(
                    id=row_id,
                    order_id=order.id,
                    order_no=order.order_no,
                    equipment_code=equipment.equipment_code if equipment else None,
                    equipment_name=equipment.name if equipment else None,
                    plan_name=plan.name if plan else None,
                    maintainer=maintainer,
                    maintenance_time=maintenance_time,
                    item_name=None,
                    item_result=None,
                    order_status=order.status,
                )
            )
            row_id += 1
    return rows


def _query_equipment_maintenance_report_orders(
    db: Session,
    *,
    keyword: str | None,
    status: str | None,
    date_from: date | None,
    date_to: date | None,
    equipment_code: str | None,
) -> list[EquipmentMaintenanceOrder]:
    query = db.query(EquipmentMaintenanceOrder).options(
        joinedload(EquipmentMaintenanceOrder.equipment),
        joinedload(EquipmentMaintenanceOrder.plan),
    )
    query = _apply_equipment_maintenance_report_filters(
        query,
        keyword=keyword,
        status=status,
        date_from=date_from,
        date_to=date_to,
        equipment_code=equipment_code,
    )
    return query.order_by(
        EquipmentMaintenanceOrder.planned_start_at.desc(),
        EquipmentMaintenanceOrder.id.desc(),
    ).all()


@router.get(
    "/equipment-maintenance",
    response_model=EquipmentMaintenanceReportListResponse,
    summary="设备保养报表",
    description=(
        "保养计划执行情况（应保养/已保养/未保养/完成率）与保养记录明细；"
        "支持关键字、状态、计划日期范围、设备编号筛选与分页。"
    ),
)
def list_equipment_maintenance_report(
    page: int = Query(1, ge=1, description="页码"),
    page_size: int = Query(10, ge=1, le=100, description="每页条数"),
    keyword: str | None = Query(None, description="工单号/设备/保养人关键字"),
    status: str | None = Query(None, description="工单状态"),
    date_from: date | None = Query(None, description="计划保养日期起（含）"),
    date_to: date | None = Query(None, description="计划保养日期止（含）"),
    equipment_code: str | None = Query(None, description="设备编号（精确匹配）"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    orders = _query_equipment_maintenance_report_orders(
        db,
        keyword=keyword,
        status=status,
        date_from=date_from,
        date_to=date_to,
        equipment_code=equipment_code,
    )
    summary = _calc_maintenance_report_summary(orders)
    all_rows = _flatten_maintenance_orders_for_report(orders)
    total = len(all_rows)
    start = (page - 1) * page_size
    items = all_rows[start : start + page_size]
    return EquipmentMaintenanceReportListResponse(
        summary=summary,
        items=items,
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get(
    "/equipment-maintenance/export",
    summary="导出设备保养报表 Excel",
    description="按当前筛选条件导出计划执行情况与保养记录明细 Excel 文件。",
)
def export_equipment_maintenance_report(
    keyword: str | None = Query(None, description="工单号/设备/保养人关键字"),
    status: str | None = Query(None, description="工单状态"),
    date_from: date | None = Query(None, description="计划保养日期起（含）"),
    date_to: date | None = Query(None, description="计划保养日期止（含）"),
    equipment_code: str | None = Query(None, description="设备编号（精确匹配）"),
    _current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    orders = _query_equipment_maintenance_report_orders(
        db,
        keyword=keyword,
        status=status,
        date_from=date_from,
        date_to=date_to,
        equipment_code=equipment_code,
    )
    summary = _calc_maintenance_report_summary(orders)
    rows = _flatten_maintenance_orders_for_report(orders)

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "计划执行情况"
    summary_ws.append(["指标", "数值"])
    summary_ws.append(["应保养", summary.due_total])
    summary_ws.append(["已保养", summary.completed])
    summary_ws.append(["未保养", summary.not_done])
    summary_ws.append(["完成率(%)", summary.completion_rate])

    detail_ws = wb.create_sheet("保养记录明细")
    detail_headers = [
        "工单号",
        "设备编号",
        "设备名称",
        "保养计划",
        "保养人",
        "保养时间",
        "保养项目",
        "结果",
        "工单状态",
    ]
    detail_ws.append(detail_headers)
    for row in rows:
        detail_ws.append(
            [
                row.order_no,
                row.equipment_code or "",
                row.equipment_name or "",
                row.plan_name or "",
                row.maintainer or "",
                row.maintenance_time.strftime("%Y-%m-%d %H:%M:%S")
                if row.maintenance_time
                else "",
                row.item_name or "",
                row.item_result or "",
                _maintenance_order_status_label(row.order_status),
            ]
        )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"equipment_maintenance_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
